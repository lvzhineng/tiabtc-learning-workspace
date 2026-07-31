"""Export the legacy BitLanglang workbook into a browser-friendly JSON snapshot."""

import hashlib
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = (
    ROOT.parent
    / "BitLanglangReview"
    / "1_2022.5~2024.6浪浪交易交割单数据3 (1).xlsx"
)
DEFAULT_TARGET = ROOT / "web" / "public" / "bitlang-trades.json"
SHEET_NAME = "时间排列+去除金额错误单子"


def number(value, default=0.0):
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return default


def iso_time(value):
    if isinstance(value, datetime):
        return value.replace(microsecond=0).isoformat() + "+08:00"
    return str(value or "").strip()


def main():
    source = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else DEFAULT_SOURCE
    target = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else DEFAULT_TARGET
    if not source.exists():
        raise FileNotFoundError(f"找不到旧交割单：{source}")

    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook[SHEET_NAME]
    rows = sheet.iter_rows(min_row=2, values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    # The worksheet contains a second analysis block with duplicate header
    # names. Preserve the first occurrence of each field so the derived
    # holding-time and amplitude columns remain available.
    header_positions = {}
    for index, header in enumerate(headers):
        if header and header not in header_positions:
            header_positions[header] = index
    trades = []

    for values in rows:
        row = {
            header: values[index]
            for header, index in header_positions.items()
            if index < len(values)
        }
        sequence = int(number(row.get("序号"), 0))
        instrument = str(row.get("交易对") or "").strip()
        direction = str(row.get("方向") or "").strip()
        entry_time = iso_time(row.get("买入时间"))
        exit_time = iso_time(row.get("卖出时间"))
        if not sequence or not instrument or direction not in {"多", "空"}:
            continue

        identity = "|".join(
            [
                str(sequence),
                instrument,
                entry_time,
                exit_time,
                direction,
                str(number(row.get("开仓均价"))),
                str(number(row.get("平仓均价"))),
                str(number(row.get("收益 (USDT)"))),
            ]
        )
        trades.append(
            {
                "id": hashlib.sha256(identity.encode("utf-8")).hexdigest(),
                "sequence": sequence,
                "instrument": instrument,
                "direction": direction,
                "leverage": number(row.get("杠杆倍数")),
                "margin": number(row.get("保证金（最大时）")),
                "entryPrice": number(row.get("开仓均价")),
                "exitPrice": number(row.get("平仓均价")),
                "returnRate": number(row.get("收益率")),
                "profit": number(row.get("收益 (USDT)")),
                "turnover": number(row.get("交易额 (USD)")),
                "size": number(row.get("持仓量（最大时）")),
                "maxPositionValue": number(row.get("持仓价值（最大时）")),
                "fee": number(row.get("手续费 (USD)")),
                "entryTime": entry_time,
                "exitTime": exit_time,
                "holdingMinutes": round(number(row.get("交易时间差（分钟）"))),
                "amplitude": number(row.get("收益率/倍数=实际振幅")),
                "sourceNote": str(row.get("备注") or "").strip(),
            }
        )

    trades.sort(key=lambda trade: trade["entryTime"])
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(
        json.dumps(
            {
                "source": source.name,
                "generatedAt": datetime.now().astimezone().isoformat(),
                "trades": trades,
            },
            ensure_ascii=False,
            separators=(",", ":"),
        ),
        encoding="utf-8",
    )
    print(f"已导出 {len(trades)} 条交割记录 -> {target}")


if __name__ == "__main__":
    main()
